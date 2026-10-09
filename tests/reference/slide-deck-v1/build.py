"""Reference solution for slide-deck-v1.

Usage: python build.py --workspace <dir>

Builds deck.pptx from the facts in source.md with python-pptx. It exists to
prove the validator accepts a correct deck. The deck is generated rather than
committed as a binary.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

SPANS = [
    ("Golden Gate", "United States", "1937", 1280),
    ("Verrazzano-Narrows", "United States", "1964", 1298),
    ("Humber", "United Kingdom", "1981", 1410),
    ("Great Belt East", "Denmark", "1998", 1624),
    ("Akashi Kaikyo", "Japan", "1998", 1991),
    ("1915 Canakkale", "Turkey", "2022", 2023),
]

BULLET_SLIDES = [
    (
        "Where it is",
        [
            "A suspension bridge across the Golden Gate strait",
            "Links San Francisco with Marin County",
            "Carries a six-lane road, a footpath and a cycle path",
        ],
        "The bridge crosses the narrow strait where San Francisco Bay meets the Pacific Ocean. "
        "It joins the city of San Francisco to Marin County in the north.",
    ),
    (
        "How it was built",
        [
            "1933: construction begins on 5 January",
            "1935: cable spinning starts between the towers",
            "1937: opens to pedestrians on 27 May and to vehicles the next day",
        ],
        "Walk through the dates in order. The bridge took four years from the first day of work to opening.",
    ),
    (
        "Size and structure",
        [
            "Main span: 1,280 m (4,200 ft)",
            "Total length: 2,737 m (8,981 ft)",
            "Towers: 227 m (746 ft) above the water",
            "Deck: 27 m wide, 67 m above the water at mid span",
        ],
        "These figures show why it was a record holder. The towers carry the two main cables and the cables carry the deck.",
    ),
    (
        "Cables and colour",
        [
            "Two main cables, each 92.4 cm thick",
            "Each cable holds 27,572 steel wires",
            "Painted International Orange so ships can see it in fog",
        ],
        "The wires were spun in place and then bound together. The orange colour was chosen to stand out against the sea and the fog.",
    ),
    (
        "Building it safely",
        [
            "Foundations built in deep, fast moving water",
            "A safety net hung under the deck during construction",
            "The net saved 19 workers; eleven workers died",
        ],
        "Be honest about the human cost. A safety net saved many lives, but a work platform fell through it in 1937.",
    ),
    (
        "Cost and a record",
        [
            "About 35 million dollars, paid by bonds repaid from tolls",
            "Longest main span in the world from 1937 until 1964",
            "1964: the Verrazzano-Narrows Bridge opens with a longer span",
        ],
        "Explain how the money was raised and how long the record lasted.",
    ),
]


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", required=True, type=Path)
    workspace = parser.parse_args().workspace

    deck = Presentation()
    title_layout, content_layout, title_only_layout = deck.slide_layouts[0], deck.slide_layouts[1], deck.slide_layouts[5]

    slide = deck.slides.add_slide(title_layout)
    slide.shapes.title.text = "The Golden Gate Bridge"
    slide.placeholders[1].text = "History and engineering"

    for title, bullets, notes in BULLET_SLIDES:
        slide = deck.slides.add_slide(content_layout)
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0]
        for line in bullets[1:]:
            frame.add_paragraph().text = line
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(24)
        add_notes(slide, notes)

    slide = deck.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "Main span compared with other bridges (m)"
    data = CategoryChartData()
    data.categories = [row[0] for row in SPANS]
    data.add_series("Main span (m)", [row[3] for row in SPANS])
    slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.7), Inches(1.6), Inches(8.6), Inches(5.4), data)
    add_notes(slide, "The Golden Gate main span of 1,280 m was a world record in 1937. Later bridges are much longer.")

    slide = deck.slides.add_slide(title_only_layout)
    slide.shapes.title.text = "Selected long suspension bridges"
    table = slide.shapes.add_table(len(SPANS) + 1, 4, Inches(0.7), Inches(1.7), Inches(8.6), Inches(3.8)).table
    for column, heading in enumerate(["Bridge", "Country", "Year opened", "Main span (m)"]):
        table.cell(0, column).text = heading
    for row_index, row in enumerate(SPANS, start=1):
        for column, value in enumerate(row):
            table.cell(row_index, column).text = str(value)
    add_notes(slide, "The same figures as the chart, with the country and the year each bridge opened.")

    slide = deck.slides.add_slide(content_layout)
    slide.shapes.title.text = "Sources"
    frame = slide.placeholders[1].text_frame
    frame.text = "Golden Gate Bridge, Wikipedia: https://en.wikipedia.org/wiki/Golden_Gate_Bridge"
    frame.add_paragraph().text = "National Park Service: https://www.nps.gov/goga/"
    frame.add_paragraph().text = "Golden Gate Bridge Highway and Transportation District: https://www.goldengate.org/"
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(18)

    deck.save(str(workspace / "deck.pptx"))


if __name__ == "__main__":
    main()
