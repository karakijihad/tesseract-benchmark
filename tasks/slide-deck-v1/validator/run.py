"""Validator for slide-deck-v1.

Usage: python validator/run.py --workspace <dir> --output <json>

Opens deck.pptx with python-pptx and checks its structure. Visual quality is
not scored here: a person judges it. Needs python-pptx.

The result file is always written. If the validator cannot run (python-pptx is
not installed) the file holds {"score": null, "max_score": N, "checks": {},
"error": "..."} and the exit status is 2. A deck that is missing, corrupt or
odd is never that error: it fails checks.
"""
from __future__ import annotations

import argparse
from collections import OrderedDict
import json
import os
from pathlib import Path
import re

HERE = Path(__file__).resolve().parent
INSTALL_HINT = "Install it with: pip install -r requirements-validators.txt"
DEFAULT_FONT_POINTS = 18.0


class EnvironmentProblem(Exception):
    """The validator itself cannot run. The message says what is missing and how to install it."""


def shape_kind(shape):
    try:
        return shape.shape_type
    except Exception:
        return None


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):  # group shape
            try:
                yield from iter_shapes(shape.shapes)
            except Exception:
                pass


def shape_chart(shape):
    try:
        return shape.chart if getattr(shape, "has_chart", False) else None
    except Exception:
        return None


def shape_table(shape):
    try:
        return shape.table if getattr(shape, "has_table", False) else None
    except Exception:
        return None


def chart_series(chart):
    """Return (category labels, list of value lists) for a chart."""
    labels: list[str] = []
    series_values: list[list[float]] = []
    try:
        for plot in chart.plots:
            try:
                labels += [str(label) for label in plot.categories]
            except Exception:
                pass
            for series in plot.series:
                try:
                    series_values.append([float(v) for v in series.values if v is not None])
                except Exception:
                    pass
    except Exception:
        pass
    return labels, series_values


def table_cells(table) -> list[str]:
    cells = []
    for row in table.rows:
        for cell in row.cells:
            cells.append(cell.text_frame.text if cell.text_frame is not None else "")
    return cells


def slide_text(slide) -> str:
    """All text a viewer can read on the slide, including tables and chart labels."""
    parts: list[str] = []
    for shape in iter_shapes(slide.shapes):
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            parts.append(shape.text_frame.text)
        table = shape_table(shape)
        if table is not None:
            parts += table_cells(table)
        chart = shape_chart(shape)
        if chart is not None:
            labels, series_values = chart_series(chart)
            parts += labels
            parts += [str(value) for values in series_values for value in values]
            try:
                if chart.has_title:
                    parts.append(chart.chart_title.text_frame.text)
            except Exception:
                pass
    return "\n".join(parts)


def notes_text(slide) -> str:
    try:
        if slide.has_notes_slide:
            return slide.notes_slide.notes_text_frame.text or ""
    except Exception:
        pass
    return ""


def slide_has_content(slide) -> bool:
    for shape in iter_shapes(slide.shapes):
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame and shape.text_frame.text.strip():
            return True
        if shape_chart(shape) is not None:
            return True
        table = shape_table(shape)
        if table is not None and any(text.strip() for text in table_cells(table)):
            return True
        if shape_kind(shape) == 13:  # picture
            return True
    return False


def font_points(shape) -> float:
    """Largest font size written on the shape's text, or the default when none is written."""
    sizes = []
    try:
        for paragraph in shape.text_frame.paragraphs:
            if paragraph.font.size is not None:
                sizes.append(paragraph.font.size.pt)
            for run in paragraph.runs:
                if run.font.size is not None:
                    sizes.append(run.font.size.pt)
    except Exception:
        pass
    return max(sizes) if sizes else DEFAULT_FONT_POINTS


def is_furniture(shape) -> bool:
    """A footer, date or slide number placeholder is not a title."""
    try:
        from pptx.enum.shapes import PP_PLACEHOLDER

        return bool(shape.is_placeholder) and shape.placeholder_format.type in (
            PP_PLACEHOLDER.DATE,
            PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.SLIDE_NUMBER,
        )
    except Exception:
        return False


def slide_title(slide) -> str:
    """The title placeholder's text, or when it has none the largest, then top-most, text on the slide."""
    try:
        title = slide.shapes.title
        if title is not None and title.has_text_frame and title.text_frame.text.strip():
            return title.text_frame.text.strip()
    except Exception:
        pass
    candidates = []
    for order, shape in enumerate(iter_shapes(slide.shapes)):
        try:
            if not (getattr(shape, "has_text_frame", False) and shape.has_text_frame) or is_furniture(shape):
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            top = shape.top if shape.top is not None else 10**12
            candidates.append((-font_points(shape), top, order, text))
        except Exception:
            continue
    if not candidates:
        return ""
    return min(candidates)[3]


def normalise(text: str) -> str:
    text = text.lower().replace("\u00a0", " ")
    text = re.sub(r"(?<=\d),(?=\d)", "", text)
    return re.sub(r"\s+", " ", text)


def squash(text: str) -> str:
    """Lower case words and domains with the punctuation between them reduced to single spaces."""
    return " " + re.sub(r"\s+", " ", re.sub(r"[^a-z0-9.]+", " ", text.lower().replace("\u00a0", " "))).strip() + " "


def contains_term(haystack: str, term: str) -> bool:
    term = normalise(term)
    if term.isdigit():
        return re.search(rf"(?<!\d){re.escape(term)}(?!\d)", haystack) is not None
    return term in haystack


def parse_number(text: str):
    cleaned = text.replace(",", "").replace("\u00a0", " ").strip()
    match = re.search(r"-?\d+(?:\.\d+)?", cleaned)
    if not match:
        return None
    try:
        return float(match.group())
    except (OverflowError, ValueError):
        return None


def same_values(found: list[float], wanted: list[float]) -> bool:
    remaining = list(found)
    for value in wanted:
        for index, candidate in enumerate(remaining):
            if abs(candidate - value) <= 1e-6 * max(1.0, abs(value)):
                del remaining[index]
                break
        else:
            return False
    return True


def find_data(presentation, columns) -> tuple[bool, bool]:
    """Return (a chart or table exists, one matches a supplied column)."""
    present = False
    matched = False
    for slide in presentation.slides:
        for shape in iter_shapes(slide.shapes):
            chart = shape_chart(shape)
            table = shape_table(shape)
            if chart is None and table is None:
                continue
            present = True
            if chart is not None:
                _, series_values = chart_series(chart)
                for values in series_values:
                    if any(same_values(values, wanted) for wanted in columns.values()):
                        matched = True
            if table is not None:
                numbers = [parse_number(text) for text in table_cells(table)]
                numbers = [number for number in numbers if number is not None]
                if any(same_values(numbers, wanted) for wanted in columns.values()):
                    matched = True
    return present, matched


def listed_sources(text: str, sources) -> list[str]:
    """The supplied sources named on the slide, by their domain or by their name."""
    lowered, squashed = text.lower(), squash(text)
    return [source["name"] for source in sources if source["domain"].lower() in lowered or squash(source["name"]) in squashed]


def score_weights() -> OrderedDict:
    scoring = json.loads((HERE.parent / "scoring.json").read_text(encoding="utf-8"))
    return OrderedDict(scoring["checks"])


def write_json(output: Path, result: dict) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(result, indent=2)
    partial = output.with_name(output.name + ".part")
    partial.write_text(text, encoding="utf-8")
    os.replace(partial, output)
    print(text)


def write_result(output: Path, weights, checks, details) -> int:
    points = {}
    passed = {}
    for name, weight in weights.items():
        value = checks.get(name, False)
        fraction = float(value) if not isinstance(value, bool) else (1.0 if value else 0.0)
        fraction = max(0.0, min(1.0, fraction))
        points[name] = round(weight * fraction, 2)
        passed[name] = fraction >= 1.0
    result = {
        "score": round(sum(points.values()), 2),
        "max_score": sum(weights.values()),
        "checks": passed,
        "points": points,
        "details": {key: str(value) for key, value in details.items() if value},
    }
    write_json(output, result)
    return 0 if result["score"] >= result["max_score"] else 1


def write_environment_error(output: Path, weights, message: str) -> int:
    write_json(output, {"score": None, "max_score": sum(weights.values()) if weights else None, "checks": {}, "error": message})
    return 2


def guarded(checks: dict, details: dict, name: str, function) -> None:
    try:
        checks[name] = function()
    except Exception as error:  # an odd deck fails the check and never crashes the validator
        checks[name] = False
        details[name] = f"{type(error).__name__}: {str(error)[:300]}"


def validate(workspace: Path, checks: dict, details: dict) -> None:
    try:
        from pptx import Presentation
    except Exception as error:
        raise EnvironmentProblem(f"python-pptx is not installed ({type(error).__name__}). {INSTALL_HINT}") from error

    expected = json.loads((HERE / "expected.json").read_text(encoding="utf-8"))
    deck_path = workspace / "deck.pptx"
    if not deck_path.is_file():
        details["opens"] = "deck.pptx is missing"
        return
    try:
        presentation = Presentation(str(deck_path))
        slides = list(presentation.slides)
    except Exception as error:  # a corrupt file is a failed check, never a crash
        details["opens"] = f"deck.pptx cannot be opened: {type(error).__name__}: {str(error)[:200]}"
        return

    count = len(slides)
    checks["opens"] = count > 0
    if count == 0:
        details["opens"] = "deck.pptx holds no slides"
    checks["slide_count"] = expected["slide_count_min"] <= count <= expected["slide_count_max"]
    details["slide_count"] = f"{count} slides"

    def title():
        first_title = slide_title(slides[0]) if slides else ""
        details["title_slide"] = f"first slide title: {first_title!r}"
        return bool(first_title)

    def empties():
        filled = [slide_has_content(slide) for slide in slides]
        details["no_empty_slides"] = f"empty slides: {[i + 1 for i, ok in enumerate(filled) if not ok]}"
        return (sum(filled) / count) if count else 0.0

    def notes():
        middle = slides[1:-1]
        minimum = expected["minimum_notes_characters"]
        with_notes = [len(notes_text(slide).strip()) >= minimum for slide in middle]
        details["speaker_notes"] = f"slides without notes: {[i + 2 for i, ok in enumerate(with_notes) if not ok]}"
        return (sum(with_notes) / len(middle)) if middle else 0.0

    def data():
        present, matched = find_data(presentation, expected["table_columns"])
        checks["data_matches_source"] = matched
        return present

    def facts():
        corpus = normalise("\n".join(slide_text(slide) + "\n" + notes_text(slide) for slide in slides))
        required = expected["required_facts"]
        found = [any(contains_term(corpus, option) for option in options) for options in required]
        details["required_facts"] = "missing: " + ", ".join(options[0] for options, ok in zip(required, found) if not ok)
        return sum(found) / len(required)

    def sources():
        last = slide_text(slides[-1]) if slides else ""
        named = listed_sources(last, expected["sources"])
        details["sources_slide"] = f"sources named on the last slide: {named}"
        return "source" in normalise(last) and len(named) >= expected["minimum_sources_on_closing_slide"]

    checks["data_matches_source"] = False
    guarded(checks, details, "title_slide", title)
    guarded(checks, details, "no_empty_slides", empties)
    guarded(checks, details, "speaker_notes", notes)
    guarded(checks, details, "chart_or_table", data)
    guarded(checks, details, "required_facts", facts)
    guarded(checks, details, "sources_slide", sources)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    workspace = args.workspace.resolve()
    weights = None
    checks: dict = {}
    details: dict = {}
    try:
        weights = score_weights()
        validate(workspace, checks, details)
    except EnvironmentProblem as problem:
        return write_environment_error(args.output, weights, str(problem))
    except Exception as error:  # a validator bug still leaves a result holding what was scored
        details["validator"] = f"validator error: {type(error).__name__}: {str(error)[:300]}"
        if weights is None:
            return write_environment_error(args.output, weights, details["validator"])
    return write_result(args.output, weights, checks, details)


if __name__ == "__main__":
    raise SystemExit(main())
